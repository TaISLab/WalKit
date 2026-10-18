"""Base of support along the walker's x axis, from the user's anthropometry and the handle height.

The allowed distance between the user and the handles is the horizontal reach of the arm holding the handlebar with
the elbow slightly flexed: shoulder, elbow and fist heights come from regressions on the user's height and age
(config/Correlationes.xlsx, one sheet per sex, one row per age bracket), and the forearm and arm angles are searched
(on a 0.1 degree grid) so that the fist is at the handlebar height with an elbow extension between 20 and 30 degrees.
The range of horizontal distances those solutions give is the base of support in x.

Pure python (numpy + openpyxl): no pandas, no ROS.
"""
import numpy as np

ELBOW_EXTENSION_DEG = (20.0, 30.0)
HEIGHT_TOLERANCE_MM = 10.0
ANGLE_STEP_DEG = 0.1


class AnthropometricTables:
    """The age-bracket regressions of one workbook, read once."""

    def __init__(self, xlsx_path):
        import openpyxl
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
        self.sheets = {}
        for name in wb.sheetnames:
            rows = list(wb[name].iter_rows(values_only=True))
            header = rows[0]
            self.sheets[name] = [dict(zip(header, r)) for r in rows[1:] if r[0] is not None]
        wb.close()

    def row(self, gender, age):
        """The bracket of that age (the nearest one if the age falls outside all of them)."""
        rows = self.sheets['AnthroprometricFemale' if gender == 'Femenino' else 'AnthroprometricMale']
        for r in rows:
            if r['Min_age'] <= age <= r['Max_age']:
                return r
        return min(rows, key=lambda r: min(abs(age - r['Min_age']), abs(age - r['Max_age'])))


def arm_geometry_mm(row, user_height_mm):
    """(shoulder height, arm length, forearm length) in mm."""
    shoulder = row['StatureX_ShoulderY_P'] * (user_height_mm - row['StatureX_ShoulderY_MeanX']) + row['StatureX_ShoulderY_MeanY']
    elbow = row['ShoulderX_ElbowY_P'] * (shoulder - row['ShoulderX_ElbowY_MeanX']) + row['ShoulderX_ElbowY_MeanY']
    fist = row['ElbowX_FistY_P'] * (elbow - row['ElbowX_FistY_MeanX']) + row['ElbowX_FistY_MeanY']
    return shoulder, shoulder - elbow, elbow - fist


def bos_x_range(row, user_height_m, handlebar_height_m):
    """(min, max) horizontal distance in m between the shoulder and the fist for the arm postures that put the fist at
    the handlebar height, or None if there is none."""
    shoulder, arm, forearm = arm_geometry_mm(row, user_height_m * 1000.0)
    shoulder_above_handlebar = shoulder - handlebar_height_m * 1000.0
    alpha = np.deg2rad(np.arange(0.0, 90.0, ANGLE_STEP_DEG))[:, None]                      # forearm angle
    extension = np.deg2rad(np.arange(ELBOW_EXTENSION_DEG[0], ELBOW_EXTENSION_DEG[1] + ANGLE_STEP_DEG / 2, ANGLE_STEP_DEG))[None, :]
    beta = alpha + extension                                                                # arm angle
    fits = np.abs(forearm * np.sin(alpha) + arm * np.sin(beta) - shoulder_above_handlebar) < HEIGHT_TOLERANCE_MM
    if not fits.any():
        return None
    reach = (forearm * np.cos(alpha) + arm * np.cos(beta))[fits] / 1000.0
    return float(reach.min()), float(reach.max())


def centered_score(value, low, high):
    """0 outside [low, high] and on its borders, 1 at its centre, linear in between ("0 in the border or beyond,
    maximum in the centre", as the index was always documented; the code ramped from 0 at one border to 1 at the
    other)."""
    if high <= low or value <= low or value >= high:
        return 0.0
    half = (high - low) / 2.0
    return 1.0 - abs(value - (low + high) / 2.0) / half
